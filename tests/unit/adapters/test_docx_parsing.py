import io
import zipfile

import pytest
from docx.document import Document as DocxDocument

from rag_anatomy.adapters.driven.parsing import DOCX, DocxParser
from rag_anatomy.domain import (
    CorruptDocumentError,
    EncryptedDocumentError,
    Page,
    UnsupportedMediaTypeError,
)
from tests.builders import make_docx


async def _text(content: bytes) -> str:
    [page] = await DocxParser().parse(content, DOCX)
    assert page.number == 1
    return page.text


async def test_paragraphs_and_tables_keep_document_order() -> None:
    def build(document: DocxDocument) -> None:
        document.add_paragraph("Before the table.")
        table = document.add_table(rows=2, cols=2)
        for row, cells in zip(table.rows, [["k", "ms"], ["10", "4.2"]], strict=True):
            for cell, text in zip(row.cells, cells, strict=True):
                cell.text = text
        document.add_paragraph("After the table.")

    assert await _text(make_docx(build)) == (
        "Before the table.\n\nk | ms\n10 | 4.2\n\nAfter the table."
    )


async def test_merged_cell_appears_once() -> None:
    def build(document: DocxDocument) -> None:
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
        table.cell(0, 2).text = "single"

    assert await _text(make_docx(build)) == "merged | single"


async def test_nested_table_text_is_included_in_its_cell() -> None:
    def build(document: DocxDocument) -> None:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "outer"
        nested = table.cell(0, 1).add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "inner"
        nested.cell(0, 1).text = "cell"

    assert await _text(make_docx(build)) == "outer | inner | cell"


async def test_headers_footers_and_empty_paragraphs_are_skipped() -> None:
    def build(document: DocxDocument) -> None:
        section = document.sections[0]
        section.header.paragraphs[0].text = "Company confidential"
        section.footer.paragraphs[0].text = "Page 1 of 9"
        document.add_paragraph("First.")
        document.add_paragraph("")
        document.add_paragraph("   ")
        document.add_paragraph("Second.")

    assert await _text(make_docx(build)) == "First.\n\nSecond."


async def test_serbian_text_survives() -> None:
    def build(document: DocxDocument) -> None:
        document.add_paragraph("Šta je đak čitao? Ćirilica: Шта је Ђак читао?")

    assert (
        await _text(make_docx(build)) == "Šta je đak čitao? Ćirilica: Шта је Ђак читао?"
    )


async def test_document_without_body_text_is_one_empty_page() -> None:
    assert await DocxParser().parse(make_docx(lambda _: None), DOCX) == [
        Page(number=1, text="")
    ]


def _zip_without_docx_parts() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "not a word document")
    return buffer.getvalue()


def _docx_with_malformed_body() -> bytes:
    source = zipfile.ZipFile(io.BytesIO(make_docx(lambda _: None)))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name in source.namelist():
            broken = name == "word/document.xml"
            archive.writestr(name, b"<w:document" if broken else source.read(name))
    return buffer.getvalue()


@pytest.mark.parametrize(
    "content",
    [
        b"just some text",
        b"PK\x03\x04garbage",
        _zip_without_docx_parts(),
        _docx_with_malformed_body(),
        b"",
    ],
    ids=["not-zip", "corrupt-zip", "zip-without-parts", "malformed-xml", "empty"],
)
async def test_unreadable_docx_is_rejected(content: bytes) -> None:
    with pytest.raises(CorruptDocumentError):
        await DocxParser().parse(content, DOCX)


async def test_password_protected_docx_is_rejected() -> None:
    compound_file = (
        bytes.fromhex("d0cf11e0a1b11ae1")
        + bytes(504)
        + "EncryptedPackage".encode("utf-16-le")
    )
    with pytest.raises(EncryptedDocumentError):
        await DocxParser().parse(compound_file, DOCX)


async def test_other_media_types_are_rejected() -> None:
    with pytest.raises(UnsupportedMediaTypeError):
        await DocxParser().parse(make_docx(lambda _: None), "application/msword")
